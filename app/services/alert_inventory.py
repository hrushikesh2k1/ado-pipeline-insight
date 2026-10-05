"""Reads an uploaded alert inventory (the alerts deployed in production) into a list of alert names.

Accepted: .xlsx, .csv, .tsv, .md and .txt. A spreadsheet or table needs a header row; the column that holds the alert
name is found by its header (or can be chosen), and one optional column (service, category ...) can group the alerts.
A Markdown or text file without a table is read as one alert per line.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import datetime, timezone
from typing import Any

from app.services.irp_format import split_table_row

MAX_BYTES = 5_000_000
MAX_ROWS = 50_000
NAME_HEADERS = ("alert name", "alertname", "alert rule name", "alert rule", "alert", "rule name", "rule", "name", "title", "display name")
CATEGORY_HEADERS = ("category", "alert type", "type", "service", "component", "application", "resource type", "area", "team", "owner")
_BULLET = re.compile(r"^\s*(?:[-*+•]|\d+[.)])\s+")


def _text(data: bytes) -> str:
    """UTF-16 only when it announces itself (without a BOM, ordinary bytes decode into nonsense); UTF-8, else Windows-1252."""
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeError:
        return data.decode("cp1252", errors="replace")


def _clean(value: Any) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def _xlsx_rows(data: bytes) -> tuple[list[list[str]], str]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [_clean(c) for c in row]
                if any(cells):
                    rows.append(cells)
                if len(rows) > MAX_ROWS:
                    raise ValueError(f"The file has more than {MAX_ROWS:,} rows.")
            if rows:
                return rows, sheet.title
    finally:
        workbook.close()
    raise ValueError("The spreadsheet is empty.")


def _delimited_rows(text: str) -> list[list[str]]:
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = []
    for row in csv.reader(io.StringIO(text), dialect):
        cells = [_clean(c) for c in row]
        if any(cells):
            rows.append(cells)
        if len(rows) > MAX_ROWS:
            raise ValueError(f"The file has more than {MAX_ROWS:,} rows.")
    return rows


def _is_divider(cells: list[str]) -> bool:
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", c.strip()) for c in cells if c.strip()) and any(cells)


def _text_rows(text: str) -> list[list[str]]:
    lines = [line for line in text.splitlines() if line.strip()]
    table = [split_table_row(line) for line in lines if line.lstrip().startswith("|")]
    if len(table) >= 2 and any(_is_divider(row) for row in table):
        return [[_clean(c) for c in row] for row in table if not _is_divider(row)]
    return [["Alert"]] + [[_clean(_BULLET.sub("", line))] for line in lines]


def _unique_headers(header: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    names = []
    for index, cell in enumerate(header, start=1):
        name = cell or f"Column {index}"
        seen[name.lower()] = seen.get(name.lower(), 0) + 1
        names.append(name if seen[name.lower()] == 1 else f"{name} ({seen[name.lower()]})")
    return names


def _pick(columns: list[str], wanted: tuple[str, ...], skip: str | None = None) -> str | None:
    lowered = {c.lower(): c for c in columns if c != skip}
    for header in wanted:
        if header in lowered:
            return lowered[header]
    return None


def parse_inventory(filename: str, data: bytes, name_column: str | None = None, category_column: str | None = None) -> dict[str, Any]:
    """The alerts in the file: {"alerts": [{"name", "category"}], "columns", "name_column", "category_column", ...}."""
    if len(data) > MAX_BYTES:
        raise ValueError(f"The file is larger than {MAX_BYTES // 1_000_000} MB.")
    extension = (filename or "").lower().rsplit(".", 1)[-1] if "." in (filename or "") else ""
    sheet = None
    if extension == "xlsx":
        rows, sheet = _xlsx_rows(data)
    elif extension in ("csv", "tsv"):
        rows = _delimited_rows(_text(data))
    elif extension in ("md", "markdown", "txt"):
        rows = _text_rows(_text(data))
    elif extension == "xls":
        raise ValueError("Old .xls files are not supported. Save the file as .xlsx or .csv and upload it again.")
    else:
        raise ValueError("Upload a .xlsx, .csv, .tsv, .md or .txt file.")
    if len(rows) < 2:
        raise ValueError("No alerts were found: the file needs a header row and at least one alert.")

    columns = _unique_headers(rows[0])
    name_col = name_column if name_column in columns else _pick(columns, NAME_HEADERS) or columns[0]
    if category_column == "":
        category_col = None
    else:
        category_col = category_column if category_column in columns and category_column != name_col else _pick(columns, CATEGORY_HEADERS, skip=name_col)
    name_index = columns.index(name_col)
    category_index = columns.index(category_col) if category_col else None

    alerts: list[dict[str, Any]] = []
    seen: set[str] = set()
    duplicates = 0
    for row in rows[1:]:
        name = row[name_index] if name_index < len(row) else ""
        if not name:
            continue
        if name.lower() in seen:
            duplicates += 1
            continue
        seen.add(name.lower())
        category = row[category_index] if category_index is not None and category_index < len(row) else ""
        alerts.append({"name": name[:300], "category": category[:100] or None})
    if not alerts:
        raise ValueError(f"No alert names were found in the column '{name_col}'.")
    return {
        "filename": filename,
        "sheet": sheet,
        "columns": columns,
        "name_column": name_col,
        "category_column": category_col,
        "alerts": alerts,
        "duplicates": duplicates,
        "uploaded_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
    }


def summarize_inventory(inventory: dict[str, Any] | None) -> dict[str, Any] | None:
    """What the page shows about the inventory (not the whole list)."""
    if not inventory:
        return None
    alerts = inventory.get("alerts") or []
    counts: dict[str, int] = {}
    for alert in alerts:
        if alert.get("category"):
            counts[alert["category"]] = counts.get(alert["category"], 0) + 1
    return {
        "filename": inventory.get("filename"),
        "sheet": inventory.get("sheet"),
        "uploaded_at": inventory.get("uploaded_at"),
        "count": len(alerts),
        "duplicates": inventory.get("duplicates", 0),
        "columns": inventory.get("columns", []),
        "name_column": inventory.get("name_column"),
        "category_column": inventory.get("category_column"),
        "categories": [{"name": n, "count": c} for n, c in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))],
        "sample": [a["name"] for a in alerts[:8]],
    }
